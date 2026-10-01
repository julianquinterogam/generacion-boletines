"""Paso 1 – Lectura de la planilla de notas (Excel) y de los logros (Word).

No modifica ningún archivo de entrada. Todo lo que no se pueda determinar
se reporta como advertencia; nunca se inventa un dato.
"""
import re
import unicodedata
from decimal import Decimal, ROUND_HALF_UP

from docx import Document
from openpyxl import load_workbook

# Orden y nombres del boletín (según BOLETIN_MODELO.xls).
# col_excel: encabezado en la planilla; clave_word: encabezado normalizado en el Word.
MATERIAS = [
    dict(col_excel="MATEM",         titulo="Matemáticas",                                 clave_word="MATEMATICAS",                         grupo=None),
    dict(col_excel="CASTELL",       titulo="Humanidades, Lengua Castellana",              clave_word="HUMANIDADES LENGUA CASTELLANO",       grupo=None),
    dict(col_excel="INGLES",        titulo="INGLES",                                      clave_word="INGLES",                              grupo="Humanidades, Lengua Extrajera"),
    dict(col_excel="FRANCES",       titulo="FRANCES",                                     clave_word="FRANCES",                             grupo="Humanidades, Lengua Extrajera"),
    dict(col_excel="LECTOE",        titulo="Lectura Crítica",                             clave_word="LECTURA CRITICA",                     grupo=None),
    dict(col_excel="NATUR.",        titulo="Ciencias Naturales y Educación Ambiental",    clave_word="CIENCIAS NATURALES Y EDUCACION AMBIENTAL", grupo=None),
    dict(col_excel="SOCIAL",        titulo="Ciencias Sociales y Democracia",              clave_word="CIENCIAS SOCIALES Y DEMOCRACIA",      grupo=None),
    dict(col_excel="TUNJA",         titulo="Tunja mi Municipio",                          clave_word="TUNJA MI",                            grupo=None),
    dict(col_excel="CAT.PAZ",       titulo="Cátedra de Paz",                              clave_word="CATEDRA DE PAZ",                      grupo=None),
    dict(col_excel="ART",           titulo="Educación Artística",                         clave_word="EDUCACION ARTISTICA",                 grupo=None),
    dict(col_excel="PROYEC.",       titulo="Proyectos lúdicos, culturales y deportivos",  clave_word="PROYECTOS LUDICOS",                   grupo=None),
    dict(col_excel="INFOR",         titulo="Tecnología e Informática",                    clave_word="TECNOLOGIA E INFORMATICA",            grupo=None),
    dict(col_excel="ETICA Y RELIG", titulo="Etica y valores - Educación Religiosa",       clave_word="ETICA Y VALORES",                     grupo=None),
    dict(col_excel="ED.FIS",        titulo="Educación Física, Recreación Y Deportes",     clave_word="EDUCACION FISICA",                    grupo=None),
]

# Encabezados del Word que agrupan materias y no llevan logros propios.
ENCABEZADOS_GRUPO = {"HUMANIDADES LENGUA EXTRANJERA"}


def _norm(texto):
    """Mayúsculas, sin tildes, sin puntuación, espacios colapsados."""
    t = unicodedata.normalize("NFD", str(texto))
    t = "".join(c for c in t if unicodedata.category(c) != "Mn")
    t = re.sub(r"[^A-Za-z0-9 ]+", " ", t.upper())
    return re.sub(r"\s+", " ", t).strip()


def desempeno(nota):
    """nota: Decimal en escala 1.0–5.0. Escala del boletín."""
    if nota >= Decimal("4.6"):
        return "SUPERIOR"
    if nota >= Decimal("4.0"):
        return "ALTO"
    if nota >= Decimal("3.5"):
        return "BÁSICO"
    return "BAJO"


def fmt(nota):
    """Formato del boletín: una cifra decimal con coma (4,9)."""
    return str(nota.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)).replace(".", ",")


# ---------------------------------------------------------------- Excel ---
def leer_notas(ruta_xlsx):
    advertencias = []
    wb = load_workbook(ruta_xlsx, data_only=True)
    ws = wb.worksheets[0]
    if len(wb.worksheets) > 1:
        advertencias.append(f"El Excel tiene {len(wb.worksheets)} hojas; solo se lee la primera ('{ws.title}').")

    titulo = " ".join(str(c.value) for row in ws.iter_rows(min_row=1, max_row=4)
                      for c in row if c.value is not None)
    grado = re.search(r"GRADO\s+(\w+)", titulo, re.I)
    periodo = re.search(r"\b(I{1,3})\s*PERIODO\s*(\d{4})", titulo, re.I)
    meta = dict(
        titulo=titulo,
        grado=grado.group(1).upper() if grado else None,
        periodo=periodo.group(1).upper() if periodo else None,
        anio=periodo.group(2) if periodo else None,
    )
    for k in ("grado", "periodo", "anio"):
        if meta[k] is None:
            advertencias.append(f"No se pudo determinar '{k}' desde el título de la planilla.")

    # Fila de encabezados
    fila_enc = next((r for r in range(1, ws.max_row + 1)
                     if _norm(ws.cell(r, 2).value or "") == "NOMBRE DEL ESTUDIANTE"), None)
    if fila_enc is None:
        raise ValueError("No se encontró la fila de encabezados ('NOMBRE DEL ESTUDIANTE').")
    cols = {str(ws.cell(fila_enc, c).value).strip(): c
            for c in range(1, ws.max_column + 1) if ws.cell(fila_enc, c).value is not None}

    faltan = [m["col_excel"] for m in MATERIAS if m["col_excel"] not in cols]
    if faltan:
        advertencias.append(f"Columnas esperadas que no están en el Excel: {faltan}")
    sobran = [h for h in cols if h not in {m["col_excel"] for m in MATERIAS}
              and h not in ("NOMBRE DEL ESTUDIANTE", "PROMEDIO", "PUESTO")]
    if sobran:
        advertencias.append(f"Columnas del Excel que no corresponden a ninguna materia conocida: {sobran}")

    estudiantes = []
    for r in range(fila_enc + 1, ws.max_row + 1):
        nombre = ws.cell(r, cols["NOMBRE DEL ESTUDIANTE"]).value
        if nombre is None or not str(nombre).strip():
            continue
        nombre = re.sub(r"\s+", " ", str(nombre)).strip()
        notas = {}
        for m in MATERIAS:
            c = cols.get(m["col_excel"])
            v = ws.cell(r, c).value if c else None
            if v is None or str(v).strip() == "":
                advertencias.append(f"{nombre}: sin nota en {m['col_excel']}.")
                notas[m["col_excel"]] = None
                continue
            try:
                n = Decimal(str(v))
            except Exception:
                advertencias.append(f"{nombre}: nota no numérica en {m['col_excel']} ({v!r}).")
                notas[m["col_excel"]] = None
                continue
            if not (Decimal(10) <= n <= Decimal(50)) or n != n.to_integral_value():
                advertencias.append(f"{nombre}: nota fuera de lo esperado en {m['col_excel']} ({v}); se esperan enteros de 10 a 50.")
                notas[m["col_excel"]] = None
                continue
            notas[m["col_excel"]] = n / 10  # 43 -> 4.3

        validas = [n for n in notas.values() if n is not None]
        promedio_exacto = (sum(validas) / len(validas)) if validas else None
        promedio_excel = ws.cell(r, cols["PROMEDIO"]).value if "PROMEDIO" in cols else None
        if promedio_exacto is not None and len(validas) == len(MATERIAS) and promedio_excel is not None:
            if abs(Decimal(str(promedio_excel)) / 10 - promedio_exacto) > Decimal("0.0001"):
                advertencias.append(f"{nombre}: el promedio del Excel ({promedio_excel}) no coincide con el recalculado.")
        estudiantes.append(dict(
            n=ws.cell(r, 1).value,
            nombre=nombre,
            notas=notas,
            promedio_exacto=promedio_exacto,
            promedio=(promedio_exacto.quantize(Decimal("0.1"), rounding=ROUND_HALF_UP)
                      if promedio_exacto is not None else None),
        ))
    return meta, estudiantes, advertencias


# ----------------------------------------------------------------- Word ---
def leer_logros(ruta_docx):
    """Devuelve ({clave_word: [logros]}, advertencias). Cada párrafo no vacío bajo
    un encabezado en negrita es un logro."""
    advertencias = []
    doc = Document(ruta_docx)
    por_materia, actual, encabezados_vistos = {}, None, []
    for p in doc.paragraphs:
        runs = [r for r in p.runs if r.text.strip()]
        texto = re.sub(r"\s+", " ", p.text).strip()
        if not texto:
            continue
        es_encabezado = bool(runs) and all(r.bold for r in runs)
        if es_encabezado:
            clave = _norm(texto)
            if clave.startswith("GRADO "):
                continue  # título del documento: "GRADO QUINTO:"
            encabezados_vistos.append(texto)
            if clave in ENCABEZADOS_GRUPO:
                actual = None
                continue
            actual = next((m["clave_word"] for m in MATERIAS
                           if clave == m["clave_word"] or clave.startswith(m["clave_word"])), None)
            if actual is None:
                advertencias.append(f"Encabezado del Word sin materia equivalente: '{texto}'.")
            else:
                por_materia.setdefault(actual, [])
        elif actual is not None:
            if re.fullmatch(r"[\W_]+", texto):
                continue  # párrafo con solo puntuación suelta
            por_materia[actual].append(texto)

    for m in MATERIAS:
        if not por_materia.get(m["clave_word"]):
            advertencias.append(f"El Word no trae logros para '{m['titulo']}'.")
    return por_materia, advertencias


# --------------------------------------------------------------- Puesto ---
def calcular_puestos(estudiantes):
    """Puesto por promedio redondeado (el que se muestra en el boletín), de mayor
    a menor. Los empatados comparten puesto y la numeración es consecutiva
    (1, 2, 2, 2, 3, 3, 4, 5)."""
    distintos = sorted({e["promedio"] for e in estudiantes if e["promedio"] is not None}, reverse=True)
    posicion = {p: i + 1 for i, p in enumerate(distintos)}
    for e in estudiantes:
        e["puesto"] = posicion.get(e["promedio"])
    return estudiantes
