"""Paso 2 – Generación del boletín (Excel editable) de un estudiante.

Replica la estructura de BOLETIN_MODELO.xls: encabezado, tabla de asignaturas con
logros y desempeño/nota a la derecha, y pie con observaciones, promedio, puesto y firmas.
Comportamiento social es SOBRESALIENTE para todos. Observaciones (con líneas para
escribir a mano) e inasistencias se dejan vacías (aún no tienen fuente).
"""
import io
import math

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, Side
from openpyxl.worksheet.pagebreak import Break

from lectura import MATERIAS, desempeno, fmt

# Director de grado según lo indicado por el colegio (clave: grado normalizado).
# Los tres grados iniciales comparten director.
DIRECTORES = {
    "PREJARDIN":  ("LAURA CAMILA JIMÉNEZ", "Directora de Grado"),
    "JARDIN":     ("LAURA CAMILA JIMÉNEZ", "Directora de Grado"),
    "TRANSICION": ("LAURA CAMILA JIMÉNEZ", "Directora de Grado"),
    "PRIMERO":    ("LORENA ROJAS HERNÁNDEZ", "Directora de Grado"),
    "SEGUNDO":    ("HILDA GAMBOA ARIZA", "Directora de Grado"),
    "TERCERO":    ("DIEGO ORLANDO AGUDELO", "Director de Grado"),
    "CUARTO":     ("ORLANDO AGUDELO HERNÁNDEZ", "Director de Grado"),
    "QUINTO":     ("ORLANDO AGUDELO HERNÁNDEZ", "Director de Grado"),
}
DIRECTORA_LICEO = ("HILDA GAMBOA ARIZA", "Directora del Liceo")
# Todos los estudiantes llevan este valor (indicado por el colegio).
COMPORTAMIENTO_SOCIAL = "SOBRESALIENTE"
PERIODO_TEXTO = {"I": "PRIMER", "II": "SEGUNDO", "III": "TERCER"}

FUENTE = "Comic Sans MS"
FUENTE_TITULO = "Bradley Hand ITC"
LADO = Side(style="thin")
MEDIO = Side(style="medium")   # marco exterior, como en la plantilla


def _clave_grado(grado):
    import unicodedata
    t = unicodedata.normalize("NFD", grado or "")
    return "".join(c for c in t if c.isalpha() and unicodedata.category(c) != "Mn").upper()


def _alto_logro(texto, caracteres_por_linea=100):
    """Alto de fila según líneas estimadas (1 línea = 18,75 ; 2 líneas = 31,5, como en el modelo).
    100 caracteres por línea: calibrado con el modelo (ahí caben líneas de hasta 107)."""
    lineas = max(1, math.ceil(len(texto) / caracteres_por_linea))
    return 12.75 * lineas + 6


def generar_boletin_xlsx(meta, estudiante, logros):
    """Devuelve (bytes_xlsx, advertencias)."""
    advertencias = []
    if meta.get("periodo") not in PERIODO_TEXTO:
        raise ValueError("No se pudo determinar el periodo desde la planilla.")
    clave = _clave_grado(meta.get("grado"))
    if clave in DIRECTORES:
        director, cargo = DIRECTORES[clave]
    else:
        director, cargo = "", "Director de Grado"
        advertencias.append(f"No hay director de grado definido para '{meta.get('grado')}'.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Boletín"
    for col, ancho in zip("ABCDEFGH", (13.1, 10.96, 10.96, 10.96, 10.96, 8.4, 13.3, 15.4)):
        ws.column_dimensions[col].width = ancho

    def f(tam, b=False, i=False, nombre=FUENTE):
        return Font(name=nombre, size=tam, bold=b, italic=i, color="FF000000")

    def escribir(fila, col, valor, fuente, h="general", v="bottom", wrap=False):
        c = ws.cell(fila, col, valor)
        c.font = fuente
        c.alignment = Alignment(horizontal=h, vertical=v, wrap_text=wrap)
        return c

    def fusionar(fila_i, col_i, fila_f, col_f):
        if (fila_i, col_i) != (fila_f, col_f):
            ws.merge_cells(start_row=fila_i, start_column=col_i, end_row=fila_f, end_column=col_f)

    # ------------------------------------------------------------ Encabezado
    escribir(1, 1, "LICEO INFANTIL SANTIAGO DE TUNJA", f(14, True, nombre=FUENTE_TITULO), "center")
    escribir(2, 1, "Licencia de funcionamiento: Resolución No. 000836 del 28 de diciembre de 2021",
             f(8, True, nombre=FUENTE_TITULO), "center")
    escribir(3, 1, f"INFORME DE EVALUACION {PERIODO_TEXTO[meta['periodo']]} PERIODO", f(11, True), "center")
    for r in (1, 2, 3, 8):
        fusionar(r, 1, r, 8)
    escribir(5, 1, "GRADO: ", f(9))
    escribir(5, 2, (meta.get("grado") or "").upper(), f(9, i=True), "left")
    escribir(5, 5, f"AÑO: {meta['anio']}", f(9))
    escribir(6, 1, "ESTUDIANTE:", f(9))
    escribir(6, 2, estudiante["nombre"], f(11, True, True))
    escribir(8, 1, " Escala valorativa de desempeño: Superior: 4.6 a 5.0      Alto: 4.0 a 4.5        "
                   "Básico: 3.5 a 3.9        Bajo: 1.0 a 3,4", f(8), "center")
    for r, h in ((1, 21), (2, 15), (3, 21.75), (4, 5.25), (5, 21), (6, 21), (7, 6), (8, 12.75), (9, 2.25), (10, 15.75)):
        ws.row_dimensions[r].height = h

    escribir(10, 1, "ASIGNATURA", f(8, True), "center")
    fusionar(10, 1, 10, 7)
    escribir(10, 8, "DESEMPEÑO", f(8, True), "center")

    # ----------------------------------------------------------------- Cuerpo
    fila = 11
    bordes_superiores = []   # (fila, col_i, col_f)
    unidades = []            # [fila_inicio, fila_fin] de cada materia o grupo (no se parten entre páginas)
    grupo_previo = None
    for m in MATERIAS:
        nota = estudiante["notas"].get(m["col_excel"])
        textos = logros.get(m["clave_word"], [])

        inicio_unidad = None
        if m["grupo"] and m["grupo"] != grupo_previo:
            inicio_unidad = fila
            bordes_superiores.append((fila, 1, 8))
            escribir(fila, 1, m["grupo"], f(8, True), "center")
            fusionar(fila, 1, fila, 7)
            ws.row_dimensions[fila].height = 17.25
            fila += 1
        grupo_previo = m["grupo"]
        if inicio_unidad is None and not m["grupo"]:
            inicio_unidad = fila

        # Fila de título de la materia + desempeño
        subtitulo = bool(m["grupo"])
        bordes_superiores.append((fila, 8, 8) if subtitulo else (fila, 1, 8))
        escribir(fila, 1, m["titulo"], f(8, True), "left" if subtitulo else "center")
        fusionar(fila, 1, fila, 7)
        escribir(fila, 8, desempeno(nota) if nota is not None else "", f(12), "center", "center", True)
        ws.row_dimensions[fila].height = 17.25
        fila += 1

        # Logros
        primera_logro = fila
        for t in textos:
            escribir(fila, 1, t, f(8), "justify", "center", True)
            fusionar(fila, 1, fila, 7)
            ws.row_dimensions[fila].height = _alto_logro(t)
            fila += 1
        # Fila en blanco de separación (en el modelo mide entre 3 y 7,5; la última es la de cierre de 1,5)
        ws.row_dimensions[fila].height = 1.5 if m is MATERIAS[-1] else 5.25
        fusionar(fila, 1, fila, 7)
        ultima = fila
        fila += 1
        if inicio_unidad is not None:
            unidades.append([inicio_unidad, ultima])
        else:
            unidades[-1][1] = ultima  # segunda materia del mismo grupo

        # Nota (celda combinada a la derecha de los logros)
        escribir(primera_logro, 8, fmt(nota) if nota is not None else "", f(12), "center", "center", True)
        fusionar(primera_logro, 8, ultima, 8)

    ultima_tabla = fila - 1

    # ---------------------------------------------------------------- Bordes
    def borde(fila_, col, **lados):
        c = ws.cell(fila_, col)
        actual = c.border
        c.border = Border(
            left=lados.get("left", actual.left), right=lados.get("right", actual.right),
            top=lados.get("top", actual.top), bottom=lados.get("bottom", actual.bottom))

    # Marco exterior de grosor medio (izquierda, derecha, arriba y cierre de la tabla)
    for r in range(1, ultima_tabla + 1):
        borde(r, 1, left=MEDIO)
        borde(r, 8, right=MEDIO)
    for c in range(1, 9):
        borde(1, c, top=MEDIO)
        borde(ultima_tabla, c, bottom=MEDIO)
        borde(2, c, bottom=LADO)    # línea bajo la licencia de funcionamiento
        borde(3, c, top=LADO)
        borde(10, c, top=LADO, bottom=LADO)
    # Línea vertical entre ASIGNATURA/logros y DESEMPEÑO/nota
    for r in range(10, ultima_tabla + 1):
        borde(r, 7, right=LADO)
        borde(r, 8, left=LADO)
    # Subrayados de GRADO y AÑO (solo bajo las celdas B y E, como en el modelo) y de ESTUDIANTE
    borde(5, 2, bottom=LADO)
    borde(5, 5, bottom=LADO)
    borde(6, 5, top=LADO)
    for c in range(2, 6):
        borde(6, c, bottom=LADO)
    for r, c1, c2 in bordes_superiores:
        for c in range(c1, c2 + 1):
            borde(r, c, top=LADO)

    # ------------------------------------------ Saltos de página sin partir materias
    def alto(r):
        return ws.row_dimensions[r].height or 15

    # Página A4 con la configuración de la plantilla: márgenes superior 0,39" e inferior 0,75",
    # escala 95 % -> (841,9 - 28,1 - 54) / 0,95 = 800 pt de filas; se deja ~5 % de holgura.
    CAPACIDAD = 760
    ALTO_PIE = 160    # observaciones, inasistencias, promedio, puesto y firmas
    y = sum(alto(r) for r in range(1, 11))
    saltos = []
    for ini_u, fin_u in unidades:
        h = sum(alto(r) for r in range(ini_u, fin_u + 1))
        if y + h > CAPACIDAD:
            saltos.append(ini_u - 1)
            y = 0
        y += h
    if y + ALTO_PIE > CAPACIDAD:
        ini_ultima = unidades[-1][0]
        saltos.append(ultima_tabla if saltos and saltos[-1] == ini_ultima - 1 else ini_ultima - 1)
    for r in saltos:
        ws.row_breaks.append(Break(id=r))
        if r != ultima_tabla:
            for c in range(1, 9):
                borde(r, c, bottom=MEDIO)      # se cierra el marco al final de la página
                borde(r + 1, c, top=MEDIO)     # y se reabre al inicio de la siguiente

    # ------------------------------------------------------------------- Pie
    # Estructura idéntica a la plantilla: fila de separación (20,25), observaciones con dos
    # líneas punteadas para escribir a mano, inasistencias/comportamiento, promedio/puesto,
    # un espacio libre para la firma (sin línea) y los nombres con su cargo.
    ws.row_dimensions[ultima_tabla + 1].height = 20.25
    fila = ultima_tabla + 2
    escribir(fila, 1, "OBSERVACIONES:", f(7))
    ws.row_dimensions[fila].height = 16.5
    ws.row_dimensions[fila + 1].height = 16.5
    PUNTEADO = Side(style="hair")
    for c in range(2, 8):
        borde(fila, c, bottom=PUNTEADO)
    for c in range(1, 8):
        borde(fila + 1, c, bottom=PUNTEADO)
    fila += 3
    escribir(fila, 1, "INASISTENCIAS:", f(7))
    escribir(fila, 5, f"COMPORTAMIENTO SOCIAL: {COMPORTAMIENTO_SOCIAL}", f(7))
    fila += 1
    escribir(fila, 1, f"PROMEDIO: {fmt(estudiante['promedio'])}", f(8))
    escribir(fila, 2, f"PUESTO: {estudiante['puesto']}", f(8), "left")
    fila += 2
    for c1, c2 in ((1, 2), (5, 6)):
        fusionar(fila, c1, fila, c2)      # espacio para la firma, sin línea
    fila += 1
    escribir(fila, 1, director, f(8, True))
    escribir(fila, 5, DIRECTORA_LICEO[0], f(8, True))
    ws.row_dimensions[fila].height = 15.75
    fila += 1
    escribir(fila, 1, cargo, f(8))
    escribir(fila, 5, DIRECTORA_LICEO[1], f(8))
    for r in range(ultima_tabla + 4, fila + 1):
        if ws.row_dimensions[r].height is None:
            ws.row_dimensions[r].height = 15

    # Las celdas vacías también usan la fuente de la plantilla (Comic Sans MS 11)
    for fila_celdas in ws.iter_rows(min_row=1, max_row=fila, min_col=1, max_col=8):
        for c in fila_celdas:
            if c.font.name not in (FUENTE, FUENTE_TITULO):
                c.font = Font(name=FUENTE, size=11, color="FF000000")

    # --------------------------------------------- Impresión A4 (como la plantilla)
    ws.page_setup.paperSize = 9
    ws.page_setup.orientation = "portrait"
    ws.page_setup.scale = 95
    ws.page_margins.left, ws.page_margins.right = 0.63, 0.24
    ws.page_margins.top, ws.page_margins.bottom = 0.39, 0.75
    ws.page_margins.header = ws.page_margins.footer = 0.51
    ws.print_options.horizontalCentered = True
    ws.print_area = f"A1:H{fila}"

    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue(), advertencias
